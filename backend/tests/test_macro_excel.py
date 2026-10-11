"""Tests for raw/ Excel workbook macroeconomic data loader."""

import pytest
import tempfile
from pathlib import Path
import pandas as pd
import openpyxl

from app.adapters.macro_excel import (
    MacroExcelDataLoader,
    MacroWorkbookError,
    get_macro_excel_loader,
)


def test_loader_workbook_and_worksheet_selection():
    loader = get_macro_excel_loader()
    assert (loader.raw_dir / "RBI_Rates.xlsx").is_file()
    assert (loader.raw_dir / "gdp_data.xlsx").is_file()
    assert (loader.raw_dir / "inflation_data.xlsx").is_file()

    # Check sheet names
    wb_rbi = openpyxl.load_workbook(loader.raw_dir / "RBI_Rates.xlsx", read_only=True)
    assert "Sheet1" in wb_rbi.sheetnames
    wb_rbi.close()

    wb_gdp = openpyxl.load_workbook(loader.raw_dir / "gdp_data.xlsx", read_only=True)
    assert "GDP 2022-23 series" in wb_gdp.sheetnames
    wb_gdp.close()

    wb_cpi = openpyxl.load_workbook(loader.raw_dir / "inflation_data.xlsx", read_only=True)
    assert "CPI 2024 series" in wb_cpi.sheetnames
    wb_cpi.close()


def test_gdp_growth_extraction():
    loader = get_macro_excel_loader()
    indicators, history = loader.load_data()
    gdp = next(i for i in indicators if i.indicator == "Real GDP Growth")

    # Correct latest observation
    assert gdp.value == 7.80
    assert gdp.unit == "%"
    assert gdp.period == "FY 2026-27 Q1"
    assert gdp.previous_value == 8.56
    assert gdp.change == -0.76
    assert "MoSPI Base 2022-23 Series" in gdp.source
    assert gdp.data_status == "official"

    # Historical series verification
    gdp_pts = history["real gdp growth"]
    assert len(gdp_pts) == 13
    assert gdp_pts[0].period == "2023-24 Q1"
    assert gdp_pts[0].value == 6.58
    assert gdp_pts[-1].period == "2026-27 Q1"
    assert gdp_pts[-1].value == 7.80

    # Chronological order verified
    years = [int(p.period.split("-")[0]) for p in gdp_pts]
    assert years == sorted(years)


def test_cpi_inflation_extraction():
    loader = get_macro_excel_loader()
    indicators, history = loader.load_data()
    cpi = next(i for i in indicators if i.indicator == "CPI Inflation")

    # Correct latest observation
    assert cpi.value == 4.82
    assert cpi.unit == "%"
    assert "August 2026" in cpi.period
    assert "Provisional" in cpi.period
    assert cpi.previous_value == 4.45
    assert cpi.change == 0.37
    assert "MoSPI CPI 2024 Series" in cpi.source
    assert cpi.data_status == "official"

    # Historical series verification
    cpi_pts = history["cpi inflation"]
    assert len(cpi_pts) == 8
    assert cpi_pts[0].period == "Jan 2026"
    assert cpi_pts[0].value == 2.74
    assert cpi_pts[-1].period == "Aug 2026"
    assert cpi_pts[-1].value == 4.82


def test_rbi_repo_rate_extraction():
    loader = get_macro_excel_loader()
    indicators, history = loader.load_data()
    repo = next(i for i in indicators if i.indicator == "Policy Repo Rate")

    # Correct latest rate change
    assert repo.value == 5.50
    assert repo.unit == "%"
    assert "07 Oct 2026" in repo.period
    assert repo.previous_value == 5.25
    assert repo.change == 0.25
    assert "RBI MPC Resolution" in repo.source
    assert repo.data_status == "official"

    # Historical step series verification
    repo_pts = history["policy repo rate"]
    assert len(repo_pts) >= 70
    assert repo_pts[-1].value == 5.50
    assert repo_pts[-2].value == 5.25

    # Check dates are chronologically sorted
    dates = [p.date for p in repo_pts if p.date]
    assert dates == sorted(dates)


def test_unemployment_completely_absent():
    loader = get_macro_excel_loader()
    indicators, history = loader.load_data()

    for ind in indicators:
        assert "unemployment" not in ind.indicator.lower()
        assert "unemployment" not in ind.source.lower()

    for key in history:
        assert "unemployment" not in key.lower()


def test_missing_workbook_raises_macro_workbook_error():
    with tempfile.TemporaryDirectory() as tmpdir:
        loader = MacroExcelDataLoader(raw_dir=tmpdir)
        with pytest.raises(MacroWorkbookError) as exc_info:
            loader.load_data(force_refresh=True)
        assert "not found" in str(exc_info.value).lower()


def test_malformed_workbook_raises_macro_workbook_error():
    with tempfile.TemporaryDirectory() as tmpdir:
        p = Path(tmpdir)
        # Create empty dummy files
        (p / "RBI_Rates.xlsx").touch()
        loader = MacroExcelDataLoader(raw_dir=tmpdir)
        with pytest.raises(MacroWorkbookError):
            loader.load_data(force_refresh=True)
