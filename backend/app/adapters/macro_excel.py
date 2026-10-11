"""Macroeconomic data loader parsing official Excel workbooks in raw/.

Reads:
- raw/RBI_Rates.xlsx (Sheet: Sheet1) -> RBI policy repo rate (%)
- raw/gdp_data.xlsx (Sheet: GDP 2022-23 series) -> Quarterly real GDP growth (%)
- raw/inflation_data.xlsx (Sheet: CPI 2024 series) -> CPI year-on-year inflation (%)

Features in-memory caching validated by file modification timestamps (mtime),
strict schema validation, chronological sorting, and explicit error handling
with zero synthetic fallbacks.
"""

from __future__ import annotations

import datetime as dt
import os
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from app.adapters.normalized import MacroHistoryPoint, MacroRecord
from app.core.logging_config import get_logger

logger = get_logger("macro_excel")


class MacroWorkbookError(Exception):
    """Raised when an Excel workbook is missing, empty, or fails schema validation."""
    pass


class MacroExcelDataLoader:
    """Loads and caches macroeconomic indicator series from raw/ Excel workbooks."""

    def __init__(self, raw_dir: Path | str | None = None) -> None:
        self.raw_dir = self._resolve_raw_dir(raw_dir)
        self._cached_indicators: list[MacroRecord] | None = None
        self._cached_history: dict[str, list[MacroHistoryPoint]] | None = None
        self._cached_metadata: dict[str, Any] = {}
        self._cached_mtimes: dict[str, float] = {}

    
    def _resolve_raw_dir(self, raw_dir: Path | str | None) -> Path:
        """Resolve the macro workbook directory across local and Vercel deployments."""
        candidates: list[Path] = []

        if raw_dir:
            candidates.append(Path(raw_dir).expanduser())

        env_dir = os.getenv("MACRO_RAW_DIR")
        if env_dir:
            candidates.append(Path(env_dir).expanduser())

        source = Path(__file__).resolve()
        backend_root = source.parents[2]
        repo_root = backend_root.parent

        candidates.extend([
            backend_root / "raw",
            repo_root / "raw",
            Path("/var/task/raw"),
            Path("/var/raw"),
        ])

        required_files = (
            "RBI_Rates.xlsx",
            "gdp_data.xlsx",
            "inflation_data.xlsx",
        )

        for candidate in candidates:
            candidate = candidate.resolve()
            if candidate.is_dir() and all(
                (candidate / filename).is_file()
                for filename in required_files
            ):
                return candidate

        return (repo_root / "raw").resolve()


    def _get_file_mtimes(self) -> dict[str, float]:
        mtimes: dict[str, float] = {}
        for fname in ["RBI_Rates.xlsx", "gdp_data.xlsx", "inflation_data.xlsx"]:
            path = self.raw_dir / fname
            if path.exists():
                mtimes[fname] = path.stat().st_mtime
            else:
                mtimes[fname] = 0.0
        return mtimes

    def _is_cache_valid(self) -> bool:
        if self._cached_indicators is None or self._cached_history is None:
            return False
        current_mtimes = self._get_file_mtimes()
        if current_mtimes != self._cached_mtimes:
            return False
        return True

    def load_data(self, force_refresh: bool = False) -> tuple[list[MacroRecord], dict[str, list[MacroHistoryPoint]]]:
        """Load and return all validated macro indicators and history series."""
        if not force_refresh and self._is_cache_valid():
            return self._cached_indicators or [], self._cached_history or {}

        logger.info("Parsing macroeconomic Excel workbooks from: %s", self.raw_dir)
        logger.info("Macro raw directory exists: %s", self.raw_dir.is_dir())
        logger.info(
            "Macro workbook availability: %s",
            {
                name: (self.raw_dir / name).is_file()
                for name in [
                    "RBI_Rates.xlsx",
                    "gdp_data.xlsx",
                    "inflation_data.xlsx",
                ]
            },
        )
        repo_record, repo_history = self._parse_rbi_repo_rates()
        gdp_record, gdp_history = self._parse_gdp_growth()
        cpi_record, cpi_history = self._parse_cpi_inflation()

        indicators: list[MacroRecord] = [
            gdp_record,
            cpi_record,
            repo_record,
        ]

        history: dict[str, list[MacroHistoryPoint]] = {
            "real gdp growth": gdp_history,
            "gdp growth": gdp_history,
            "quarterly real gdp growth": gdp_history,
            "gdp": gdp_history,
            "cpi inflation": cpi_history,
            "inflation": cpi_history,
            "cpi": cpi_history,
            "cpi combined inflation": cpi_history,
            "policy repo rate": repo_history,
            "repo rate": repo_history,
            "rbi policy repo rate": repo_history,
            "interest rate": repo_history,
        }

        self._cached_indicators = indicators
        self._cached_history = history
        self._cached_mtimes = self._get_file_mtimes()
        return indicators, history

    def _parse_rbi_repo_rates(self) -> tuple[MacroRecord, list[MacroHistoryPoint]]:
        """Parse raw/RBI_Rates.xlsx sheet 'Sheet1'."""
        path = self.raw_dir / "RBI_Rates.xlsx"
        if not path.is_file():
            raise MacroWorkbookError(f"RBI rates workbook not found at: {path}")

        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            if "Sheet1" not in wb.sheetnames:
                raise MacroWorkbookError(
                    f"RBI_Rates.xlsx missing required sheet 'Sheet1'. Available: {wb.sheetnames}"
                )
            wb.close()
        except Exception as e:
            if isinstance(e, MacroWorkbookError):
                raise
            raise MacroWorkbookError(f"Failed to inspect RBI_Rates.xlsx: {e}") from e

        try:
            df_raw = pd.read_excel(path, sheet_name="Sheet1", header=None)
        except Exception as e:
            raise MacroWorkbookError(f"Error reading RBI_Rates.xlsx Sheet1: {e}") from e

        if df_raw.empty:
            raise MacroWorkbookError("RBI_Rates.xlsx Sheet1 is empty.")

        header_idx = None
        for i in range(min(15, len(df_raw))):
            row_vals = [str(x).lower() for x in df_raw.iloc[i].values if pd.notna(x)]
            if any("effective date" in v for v in row_vals):
                header_idx = i
                break

        if header_idx is None:
            raise MacroWorkbookError(
                "Could not locate 'Effective Date' header row in RBI_Rates.xlsx Sheet1."
            )

        df = pd.read_excel(path, sheet_name="Sheet1", skiprows=header_idx)
        date_cols = [c for c in df.columns if "effective date" in str(c).lower()]
        repo_cols = [c for c in df.columns if "repo rate" in str(c).lower()]

        if not date_cols or not repo_cols:
            raise MacroWorkbookError(
                f"Missing required columns in RBI_Rates.xlsx. Found: {list(df.columns)}"
            )

        date_col = date_cols[0]
        repo_col = repo_cols[0]

        records: list[dict[str, Any]] = []
        for _, row in df.iterrows():
            d = row[date_col]
            r = row[repo_col]
            if pd.isna(d) or pd.isna(r):
                continue

            # Skip footnote rows
            if isinstance(d, str):
                s_stripped = d.strip()
                if (
                    s_stripped.startswith("1.")
                    or s_stripped.startswith("2.")
                    or "update" in s_stripped.lower()
                    or "notes" in s_stripped.lower()
                ):
                    continue

            # Dashes, blanks, non-numeric mean no change in policy rate
            r_str = str(r).strip()
            if r_str in ["-", "", "nan", "--"]:
                continue

            try:
                rate_val = float(r)
                parsed_date = pd.to_datetime(d)
                records.append({
                    "date": parsed_date,
                    "rate": rate_val,
                })
            except Exception:
                continue

        if not records:
            raise MacroWorkbookError(
                "No valid numeric repo rate change records found in RBI_Rates.xlsx Sheet1."
            )

        # Sort chronologically (oldest to newest)
        records.sort(key=lambda x: x["date"])

        history_points: list[MacroHistoryPoint] = []
        for r in records:
            d_obj = r["date"]
            history_points.append(
                MacroHistoryPoint(
                    period=d_obj.strftime("%d %b %Y"),
                    value=round(r["rate"], 2),
                    date=d_obj.strftime("%Y-%m-%d"),
                )
            )

        latest = records[-1]
        prior = records[-2] if len(records) > 1 else None
        prev_val = round(prior["rate"], 2) if prior else None
        change = round(latest["rate"] - prev_val, 2) if prev_val is not None else 0.0

        latest_dt = latest["date"]
        period_str = f"Effective {latest_dt.strftime('%d %b %Y')}"

        record = MacroRecord(
            indicator="Policy Repo Rate",
            value=round(latest["rate"], 2),
            unit="%",
            period=period_str,
            previous_value=prev_val,
            change=change,
            source="raw/RBI_Rates.xlsx (RBI MPC Resolution)",
            provider="excel_macro",
            data_status="official",
            retrieved_at=dt.datetime.now(dt.timezone.utc),
        )

        return record, history_points

    def _parse_gdp_growth(self) -> tuple[MacroRecord, list[MacroHistoryPoint]]:
        """Parse raw/gdp_data.xlsx sheet 'GDP 2022-23 series'."""
        path = self.raw_dir / "gdp_data.xlsx"
        if not path.is_file():
            raise MacroWorkbookError(f"GDP workbook not found at: {path}")

        sheet_name = "GDP 2022-23 series"
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            if sheet_name not in wb.sheetnames:
                raise MacroWorkbookError(
                    f"gdp_data.xlsx missing required sheet '{sheet_name}'. Available: {wb.sheetnames}"
                )
            wb.close()
        except Exception as e:
            if isinstance(e, MacroWorkbookError):
                raise
            raise MacroWorkbookError(f"Failed to inspect gdp_data.xlsx: {e}") from e

        try:
            df_raw = pd.read_excel(path, sheet_name=sheet_name, header=None)
        except Exception as e:
            raise MacroWorkbookError(f"Error reading gdp_data.xlsx '{sheet_name}': {e}") from e

        if df_raw.empty:
            raise MacroWorkbookError(f"gdp_data.xlsx '{sheet_name}' is empty.")

        header_idx = None
        for i in range(min(15, len(df_raw))):
            row_vals = [str(x).lower() for x in df_raw.iloc[i].values if pd.notna(x)]
            if any("financial year" in v for v in row_vals):
                header_idx = i
                break

        if header_idx is None:
            raise MacroWorkbookError(
                f"Could not locate 'Financial year' header in gdp_data.xlsx '{sheet_name}'."
            )

        df = pd.read_excel(path, sheet_name=sheet_name, skiprows=header_idx)
        fy_cols = [c for c in df.columns if "financial year" in str(c).lower()]
        q_cols = [c for c in df.columns if "quarter" in str(c).lower()]
        gdp_cols = [c for c in df.columns if "real gdp" in str(c).lower()]

        if not fy_cols or not q_cols or not gdp_cols:
            raise MacroWorkbookError(
                f"Missing required columns (Financial year, Quarter, Real GDP) in '{sheet_name}'. Found: {list(df.columns)}"
            )

        fy_col = fy_cols[0]
        q_col = q_cols[0]
        gdp_col = gdp_cols[0]

        records: list[dict[str, Any]] = []
        for _, row in df.iterrows():
            fy = row[fy_col]
            q = row[q_col]
            val = row[gdp_col]
            if pd.isna(fy) or pd.isna(q) or pd.isna(val):
                continue

            fy_str = str(fy).strip()
            q_str = str(q).strip().upper()
            if not q_str.startswith("Q") or len(q_str) < 2 or not q_str[1].isdigit():
                continue

            try:
                gdp_val = float(val)
                start_year = int(fy_str.split("-")[0])
                q_num = int(q_str[1])
                # Approximate quarter end date for calendar plotting
                # Q1: Apr-Jun, Q2: Jul-Sep, Q3: Oct-Dec, Q4: Jan-Mar
                calendar_year = start_year if q_num in [1, 2, 3] else start_year + 1
                month_day_map = {1: (6, 30), 2: (9, 30), 3: (12, 31), 4: (3, 31)}
                m, d = month_day_map[q_num]
                q_date = f"{calendar_year}-{m:02d}-{d:02d}"

                records.append({
                    "fy": fy_str,
                    "quarter": q_str,
                    "period": f"{fy_str} {q_str}",
                    "value": round(gdp_val, 2),
                    "date": q_date,
                    "sort_key": (start_year, q_num),
                })
            except Exception:
                continue

        if not records:
            raise MacroWorkbookError(
                f"No valid real GDP observations extracted from '{sheet_name}'."
            )

        # Sort chronologically
        records.sort(key=lambda x: x["sort_key"])

        history_points: list[MacroHistoryPoint] = []
        for r in records:
            history_points.append(
                MacroHistoryPoint(
                    period=r["period"],
                    value=r["value"],
                    date=r["date"],
                )
            )

        latest = records[-1]
        prior = records[-2] if len(records) > 1 else None
        prev_val = prior["value"] if prior else None
        change = round(latest["value"] - prev_val, 2) if prev_val is not None else 0.0

        record = MacroRecord(
            indicator="Real GDP Growth",
            value=latest["value"],
            unit="%",
            period=f"FY {latest['period']}",
            previous_value=prev_val,
            change=change,
            source="raw/gdp_data.xlsx (MoSPI Base 2022-23 Series)",
            provider="excel_macro",
            data_status="official",
            retrieved_at=dt.datetime.now(dt.timezone.utc),
        )

        return record, history_points

    def _parse_cpi_inflation(self) -> tuple[MacroRecord, list[MacroHistoryPoint]]:
        """Parse raw/inflation_data.xlsx sheet 'CPI 2024 series'."""
        path = self.raw_dir / "inflation_data.xlsx"
        if not path.is_file():
            raise MacroWorkbookError(f"Inflation workbook not found at: {path}")

        sheet_name = "CPI 2024 series"
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            if sheet_name not in wb.sheetnames:
                raise MacroWorkbookError(
                    f"inflation_data.xlsx missing required sheet '{sheet_name}'. Available: {wb.sheetnames}"
                )
            wb.close()
        except Exception as e:
            if isinstance(e, MacroWorkbookError):
                raise
            raise MacroWorkbookError(f"Failed to inspect inflation_data.xlsx: {e}") from e

        try:
            df_raw = pd.read_excel(path, sheet_name=sheet_name, header=None)
        except Exception as e:
            raise MacroWorkbookError(f"Error reading inflation_data.xlsx '{sheet_name}': {e}") from e

        if df_raw.empty:
            raise MacroWorkbookError(f"inflation_data.xlsx '{sheet_name}' is empty.")

        header_idx = None
        for i in range(min(15, len(df_raw))):
            row_vals = [str(x).lower() for x in df_raw.iloc[i].values if pd.notna(x)]
            if any("month" in v for v in row_vals):
                header_idx = i
                break

        if header_idx is None:
            raise MacroWorkbookError(
                f"Could not locate 'Month' header row in inflation_data.xlsx '{sheet_name}'."
            )

        df = pd.read_excel(path, sheet_name=sheet_name, skiprows=header_idx)
        m_cols = [c for c in df.columns if "month" in str(c).lower()]
        inf_cols = [c for c in df.columns if "inflation combined" in str(c).lower()]
        inf_r_cols = [c for c in df.columns if "inflation rural" in str(c).lower()]
        inf_u_cols = [c for c in df.columns if "inflation urban" in str(c).lower()]
        st_cols = [c for c in df.columns if "status" in str(c).lower()]

        if not m_cols or not inf_cols:
            raise MacroWorkbookError(
                f"Missing required columns (Month, Inflation combined (%)) in '{sheet_name}'. Found: {list(df.columns)}"
            )

        m_col = m_cols[0]
        inf_col = inf_cols[0]
        inf_r_col = inf_r_cols[0] if inf_r_cols else None
        inf_u_col = inf_u_cols[0] if inf_u_cols else None
        st_col = st_cols[0] if st_cols else None

        records: list[dict[str, Any]] = []
        for _, row in df.iterrows():
            m = row[m_col]
            inf = row[inf_col]
            if pd.isna(m) or pd.isna(inf):
                continue

            try:
                dt_val = pd.to_datetime(m)
                inf_val = float(inf)
                r_val = float(row[inf_r_col]) if (inf_r_col and pd.notna(row[inf_r_col])) else None
                u_val = float(row[inf_u_col]) if (inf_u_col and pd.notna(row[inf_u_col])) else None
                status = str(row[st_col]).strip() if (st_col and pd.notna(row[st_col])) else None

                records.append({
                    "date": dt_val,
                    "period": dt_val.strftime("%b %Y"),
                    "value": round(inf_val, 2),
                    "rural_value": round(r_val, 2) if r_val is not None else None,
                    "urban_value": round(u_val, 2) if u_val is not None else None,
                    "status": status,
                })
            except Exception:
                continue

        if not records:
            raise MacroWorkbookError(
                f"No valid published year-on-year inflation records found in '{sheet_name}'."
            )

        # Sort chronologically (oldest to newest)
        records.sort(key=lambda x: x["date"])

        history_points: list[MacroHistoryPoint] = []
        for r in records:
            history_points.append(
                MacroHistoryPoint(
                    period=r["period"],
                    value=r["value"],
                    date=r["date"].strftime("%Y-%m-%d"),
                )
            )

        latest = records[-1]
        prior = records[-2] if len(records) > 1 else None
        prev_val = prior["value"] if prior else None
        change = round(latest["value"] - prev_val, 2) if prev_val is not None else 0.0

        status_suffix = f" ({latest['status']})" if latest.get("status") else ""
        period_str = f"{latest['date'].strftime('%B %Y')}{status_suffix}"

        record = MacroRecord(
            indicator="CPI Inflation",
            value=latest["value"],
            unit="%",
            period=period_str,
            previous_value=prev_val,
            change=change,
            source="raw/inflation_data.xlsx (MoSPI CPI 2024 Series)",
            provider="excel_macro",
            data_status="official",
            retrieved_at=dt.datetime.now(dt.timezone.utc),
        )

        return record, history_points


# Singleton instance
_excel_loader: MacroExcelDataLoader | None = None


def get_macro_excel_loader() -> MacroExcelDataLoader:
    global _excel_loader
    if _excel_loader is None:
        _excel_loader = MacroExcelDataLoader()
    return _excel_loader
